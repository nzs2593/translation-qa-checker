"""Convert the supplied glossary workbook to a browser-loadable data module."""

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "5820713_glossary-0-3.xlsx"
OUTPUT = ROOT / "src" / "glossary-data.js"
LANGUAGES = ("en-GB", "az", "da", "mn", "pt-BR", "ru", "es", "es-AR", "uk", "uz", "uzb")


def main():
    worksheet = load_workbook(SOURCE, read_only=True, data_only=True).active
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    columns = {name: index for index, name in enumerate(headers) if name}
    entries = []
    for row_number, row in enumerate(rows, start=2):
        row = tuple(row) + (None,) * (len(headers) - len(row))
        terms = {}
        notes = {}
        for language in LANGUAGES:
            term = row[columns[f"Term [{language}]"]]
            note = row[columns[f"Note [{language}]"]]
            if term is not None and str(term).strip():
                terms[language] = str(term).strip()
            if note is not None and str(note).strip():
                notes[language] = str(note).strip()
        if terms.get("en-GB"):
            entry = {"id": row_number, "terms": terms}
            if notes:
                entry["notes"] = notes
            entries.append(entry)
    content = "// Generated from data/5820713_glossary-0-3.xlsx. Edit the workbook and rerun scripts/build-glossary-data.py to refresh defaults.\n"
    content += "export const GLOSSARY_DEFAULTS = " + json.dumps(entries, ensure_ascii=False, separators=(",", ":")) + ";\n"
    OUTPUT.write_text(content, encoding="utf-8")
    print(f"Wrote {len(entries)} glossary rows to {OUTPUT}")


if __name__ == "__main__":
    main()
